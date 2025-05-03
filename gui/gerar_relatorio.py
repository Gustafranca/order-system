import os
import sys
from tkinter import messagebox
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from docx import Document
from docx.shared import Pt

def gerar_relatorio(nome_comprador, itens, fretes):
    if getattr(sys, 'frozen', False):
        pasta_base = os.path.dirname(sys.executable)
    else:
        pasta_base = os.path.dirname(os.path.abspath(__file__))

    # Pastas de saída
    pasta_pdf = os.path.join(pasta_base, "relatorio", "pdf")
    pasta_docx = os.path.join(pasta_base, "relatorio", "docx")
    os.makedirs(pasta_pdf, exist_ok=True)
    os.makedirs(pasta_docx, exist_ok=True)

    nome_arquivo_base = f"pedido_{nome_comprador.replace(' ', '_')}"

    # ========== PDF ==========
    nome_arquivo_pdf = os.path.join(pasta_pdf, nome_arquivo_base + ".pdf")
    c = canvas.Canvas(nome_arquivo_pdf, pagesize=A4)
    largura, altura = A4

    y = altura - 2*cm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(2*cm, y, f"Relatório de Pedido - {nome_comprador}")

    y -= 1.5*cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2*cm, y, "Itens do Pedido:")

    y -= 1*cm
    c.setFont("Helvetica", 11)
    total = 0
    c.drawString(2*cm, y, "Qtd")
    c.drawString(4*cm, y, "Item")
    c.drawString(12*cm, y, "Valor Unit.")
    c.drawString(15*cm, y, "Subtotal")

    y -= 0.5*cm
    for item in itens:
        if y < 3*cm:
            c.showPage()
            y = altura - 2*cm
        nome = item['nome']
        qtd = item['qtd']
        valor = item['valor_unit']
        subtotal = qtd * valor
        total += subtotal
        c.drawString(2*cm, y, str(qtd))
        c.drawString(4*cm, y, nome[:40])
        c.drawString(12*cm, y, f"R$ {valor:.2f}")
        c.drawString(15*cm, y, f"R$ {subtotal:.2f}")
        y -= 0.5*cm

    y -= 1*cm
    c.setFont("Helvetica-Bold", 11)
    c.drawString(12*cm, y, f"Subtotal: R$ {total:.2f}")
    y -= 0.5*cm

    frete_total = sum(fretes)
    c.drawString(12*cm, y, "Fretes:")
    y -= 0.5*cm
    for frete in fretes:
        c.drawString(13*cm, y, f"R$ {frete:.2f}")
        y -= 0.4*cm
        if y < 3*cm:
            c.showPage()
            y = altura - 2*cm

    y -= 0.2*cm
    c.drawString(12*cm, y, f"Total: R$ {total + frete_total:.2f}")
    c.save()

    # ========== DOCX ==========
    nome_arquivo_docx = os.path.join(pasta_docx, nome_arquivo_base + ".docx")
    doc = Document()
    doc.add_heading(f'Relatório de Pedido - {nome_comprador}', level=1)

    doc.add_heading('Itens do Pedido:', level=2)
    tabela = doc.add_table(rows=1, cols=4)
    tabela.style = 'Table Grid'
    hdr_cells = tabela.rows[0].cells
    hdr_cells[0].text = 'Qtd'
    hdr_cells[1].text = 'Item'
    hdr_cells[2].text = 'Valor Unit.'
    hdr_cells[3].text = 'Subtotal'

    total = 0
    for item in itens:
        nome = item['nome']
        qtd = item['qtd']
        valor = item['valor_unit']
        subtotal = qtd * valor
        total += subtotal
        row_cells = tabela.add_row().cells
        row_cells[0].text = str(qtd)
        row_cells[1].text = nome
        row_cells[2].text = f"R$ {valor:.2f}"
        row_cells[3].text = f"R$ {subtotal:.2f}"

    doc.add_paragraph(f"\nSubtotal: R$ {total:.2f}")
    doc.add_paragraph("Fretes:")
    for frete in fretes:
        doc.add_paragraph(f"  R$ {frete:.2f}")
    doc.add_paragraph(f"Total: R$ {total + sum(fretes):.2f}")

    doc.save(nome_arquivo_docx)

    # Mensagem final
    messagebox.showinfo("Relatório", f"Arquivos gerados com sucesso:\n\nPDF: {nome_arquivo_pdf}\nDOCX: {nome_arquivo_docx}")



